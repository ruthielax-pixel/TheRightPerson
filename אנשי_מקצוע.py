import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl
import os
import webbrowser

EXCEL_FILE = os.path.join(os.path.dirname(__file__), "אנשי מקצוע מומלצים.xlsx")
COLUMNS = ["שם", "כתובת קליניקה", "נייד", "אימייל", "אתר", "התמחות", "גישה", "סרטון הכרות"]
COL_WIDTHS = [160, 200, 110, 200, 180, 110, 180, 120]

BG = "#f5f5f0"
HEADER_BG = "#4a7c59"
HEADER_FG = "white"
ACCENT = "#4a7c59"
ROW_ODD = "#ffffff"
ROW_EVEN = "#eef3ee"
SELECT_BG = "#b8d4bf"
FONT_MAIN = ("Arial", 11)
FONT_BOLD = ("Arial", 11, "bold")
FONT_TITLE = ("Arial", 16, "bold")


def load_data():
    wb = openpyxl.load_workbook(EXCEL_FILE)
    data = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(cell is not None for cell in row):
                cleaned = []
                for i, val in enumerate(row):
                    if i == 2 and isinstance(val, float):
                        val = str(int(val))
                    cleaned.append(val if val is not None else "")
                while len(cleaned) < len(COLUMNS):
                    cleaned.append("")
                rows.append(cleaned[:len(COLUMNS)])
        data[sheet_name] = rows
    return data


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("אנשי מקצוע מומלצים")
        self.geometry("1100x680")
        self.minsize(800, 500)
        self.configure(bg=BG)
        self.option_add("*Font", FONT_MAIN)

        self.all_data = load_data()
        self.sheet_names = list(self.all_data.keys())
        self.current_sheet = tk.StringVar(value=self.sheet_names[0])
        self.search_var = tk.StringVar()
        self.filter_hitmaheut = tk.StringVar(value="הכל")
        self.filter_gisha = tk.StringVar(value="הכל")

        self.search_var.trace_add("write", lambda *_: self.apply_filters())
        self.filter_hitmaheut.trace_add("write", lambda *_: self.apply_filters())
        self.filter_gisha.trace_add("write", lambda *_: self.apply_filters())

        self._build_ui()
        self.switch_sheet()

    def _build_ui(self):
        # Title bar
        title_frame = tk.Frame(self, bg=HEADER_BG, pady=10)
        title_frame.pack(fill="x")
        tk.Label(title_frame, text="אנשי מקצוע מומלצים", bg=HEADER_BG, fg=HEADER_FG,
                 font=FONT_TITLE).pack()

        # Tab bar for sheets
        tab_frame = tk.Frame(self, bg=ACCENT, pady=4)
        tab_frame.pack(fill="x")
        self.tab_buttons = {}
        for sheet in self.sheet_names:
            btn = tk.Button(
                tab_frame, text=sheet, font=FONT_MAIN,
                bg=ACCENT, fg="white", relief="flat", padx=14, pady=6,
                activebackground="#5a9c6e", activeforeground="white",
                cursor="hand2",
                command=lambda s=sheet: self._on_tab(s)
            )
            btn.pack(side="right", padx=2)
            self.tab_buttons[sheet] = btn

        # Search & filter bar
        filter_frame = tk.Frame(self, bg=BG, pady=8, padx=12)
        filter_frame.pack(fill="x")

        tk.Label(filter_frame, text="חיפוש:", bg=BG, font=FONT_BOLD).pack(side="right")
        search_entry = tk.Entry(filter_frame, textvariable=self.search_var,
                                width=28, font=FONT_MAIN, justify="right")
        search_entry.pack(side="right", padx=(0, 6))

        tk.Label(filter_frame, text="   התמחות:", bg=BG, font=FONT_BOLD).pack(side="right")
        self.cb_hitmaheut = ttk.Combobox(filter_frame, textvariable=self.filter_hitmaheut,
                                          width=18, font=FONT_MAIN, state="readonly", justify="right")
        self.cb_hitmaheut.pack(side="right", padx=(0, 6))

        tk.Label(filter_frame, text="   גישה:", bg=BG, font=FONT_BOLD).pack(side="right")
        self.cb_gisha = ttk.Combobox(filter_frame, textvariable=self.filter_gisha,
                                      width=18, font=FONT_MAIN, state="readonly", justify="right")
        self.cb_gisha.pack(side="right", padx=(0, 6))

        btn_clear = tk.Button(filter_frame, text="נקה סינון", font=FONT_MAIN,
                              bg="#c0392b", fg="white", relief="flat", padx=10,
                              cursor="hand2", command=self.clear_filters)
        btn_clear.pack(side="left", padx=6)

        self.count_label = tk.Label(filter_frame, text="", bg=BG, font=FONT_MAIN, fg="#555")
        self.count_label.pack(side="left", padx=10)

        # Results table
        table_frame = tk.Frame(self, bg=BG)
        table_frame.pack(fill="both", expand=True, padx=10, pady=(0, 4))

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Treeview", font=FONT_MAIN, rowheight=28,
                         background=ROW_ODD, fieldbackground=ROW_ODD)
        style.configure("Treeview.Heading", font=FONT_BOLD,
                         background=HEADER_BG, foreground=HEADER_FG)
        style.map("Treeview", background=[("selected", SELECT_BG)],
                  foreground=[("selected", "#000")])

        cols = COLUMNS
        self.tree = ttk.Treeview(table_frame, columns=cols, show="headings",
                                  selectmode="browse")
        for col, width in zip(COLUMNS, COL_WIDTHS):
            self.tree.heading(col, text=col, anchor="e")
            self.tree.column(col, width=width, anchor="e", stretch=True)

        scrollbar_y = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        scrollbar_x = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=scrollbar_y.set, xscrollcommand=scrollbar_x.set)

        scrollbar_y.pack(side="left", fill="y")
        scrollbar_x.pack(side="bottom", fill="x")
        self.tree.pack(fill="both", expand=True)

        self.tree.tag_configure("odd", background=ROW_ODD)
        self.tree.tag_configure("even", background=ROW_EVEN)

        self.tree.bind("<<TreeviewSelect>>", self._on_select)
        self.tree.bind("<Double-1>", self._on_double_click)

        # Detail panel
        detail_outer = tk.Frame(self, bg="#dde8de", pady=8, padx=12)
        detail_outer.pack(fill="x", padx=10, pady=(0, 8))
        tk.Label(detail_outer, text="פרטים:", bg="#dde8de", font=FONT_BOLD).pack(anchor="e")
        self.detail_label = tk.Label(detail_outer, text="לחץ על שורה לצפייה בפרטים",
                                      bg="#dde8de", font=FONT_MAIN, justify="right",
                                      anchor="e", wraplength=900)
        self.detail_label.pack(anchor="e")

        self.open_link_btn = tk.Button(detail_outer, text="פתח אתר", font=FONT_MAIN,
                                        bg=ACCENT, fg="white", relief="flat", padx=10,
                                        cursor="hand2", command=self._open_website)
        self.open_link_btn.pack(anchor="w")
        self.open_link_btn.pack_forget()
        self._selected_website = ""

    def _on_tab(self, sheet_name):
        self.current_sheet.set(sheet_name)
        self.switch_sheet()

    def switch_sheet(self):
        sheet = self.current_sheet.get()
        for name, btn in self.tab_buttons.items():
            if name == sheet:
                btn.configure(bg="#2d5a3e", relief="sunken")
            else:
                btn.configure(bg=ACCENT, relief="flat")

        rows = self.all_data[sheet]
        hitmaheut_vals = sorted(set(r[5] for r in rows if r[5]))
        gisha_vals = sorted(set(r[6] for r in rows if r[6]))
        self.cb_hitmaheut["values"] = ["הכל"] + hitmaheut_vals
        self.cb_gisha["values"] = ["הכל"] + gisha_vals
        self.filter_hitmaheut.set("הכל")
        self.filter_gisha.set("הכל")
        self.search_var.set("")
        self.apply_filters()

    def apply_filters(self):
        sheet = self.current_sheet.get()
        rows = self.all_data[sheet]
        query = self.search_var.get().strip().lower()
        hitmaheut = self.filter_hitmaheut.get()
        gisha = self.filter_gisha.get()

        results = []
        for row in rows:
            if hitmaheut != "הכל" and row[5] != hitmaheut:
                continue
            if gisha != "הכל" and row[6] != gisha:
                continue
            if query and not any(query in str(v).lower() for v in row):
                continue
            results.append(row)

        self._populate_table(results)

    def _populate_table(self, results):
        self.tree.delete(*self.tree.get_children())
        for i, row in enumerate(results):
            tag = "odd" if i % 2 == 0 else "even"
            self.tree.insert("", "end", values=row, tags=(tag,))
        count = len(results)
        self.count_label.configure(text=f"נמצאו {count} תוצאות")
        self.detail_label.configure(text="לחץ על שורה לצפייה בפרטים")
        self.open_link_btn.pack_forget()

    def _on_select(self, event):
        selected = self.tree.selection()
        if not selected:
            return
        values = self.tree.item(selected[0], "values")
        labels = COLUMNS
        text = "   |   ".join(f"{labels[i]}: {values[i]}" for i in range(len(labels)) if values[i])
        self.detail_label.configure(text=text)
        self._selected_website = values[4] if len(values) > 4 else ""
        if self._selected_website:
            self.open_link_btn.pack(anchor="w")
        else:
            self.open_link_btn.pack_forget()

    def _on_double_click(self, event):
        selected = self.tree.selection()
        if not selected:
            return
        values = self.tree.item(selected[0], "values")
        win = tk.Toplevel(self)
        win.title(f"פרטי איש מקצוע: {values[0]}")
        win.geometry("500x380")
        win.configure(bg=BG)
        win.grab_set()

        tk.Label(win, text=values[0], bg=BG, font=FONT_TITLE, fg=ACCENT).pack(pady=(16, 8))
        frame = tk.Frame(win, bg=BG)
        frame.pack(fill="both", expand=True, padx=20)

        for i, (label, val) in enumerate(zip(COLUMNS, values)):
            if not val:
                continue
            row_frame = tk.Frame(frame, bg=BG, pady=3)
            row_frame.pack(fill="x")
            tk.Label(row_frame, text=f"{label}:", bg=BG, font=FONT_BOLD,
                     fg=ACCENT, width=14, anchor="e").pack(side="right")
            tk.Label(row_frame, text=val, bg=BG, font=FONT_MAIN,
                     anchor="e", justify="right", wraplength=300).pack(side="right")

        if values[4]:
            tk.Button(win, text=f"פתח אתר: {values[4]}", font=FONT_MAIN,
                      bg=ACCENT, fg="white", relief="flat", padx=10, pady=4,
                      cursor="hand2",
                      command=lambda: webbrowser.open(values[4])).pack(pady=10)
        tk.Button(win, text="סגור", font=FONT_MAIN, bg="#777", fg="white",
                  relief="flat", padx=16, command=win.destroy).pack(pady=4)

    def _open_website(self):
        url = self._selected_website
        if url and not url.startswith("http"):
            url = "https://" + url
        if url:
            webbrowser.open(url)

    def clear_filters(self):
        self.search_var.set("")
        self.filter_hitmaheut.set("הכל")
        self.filter_gisha.set("הכל")


if __name__ == "__main__":
    app = App()
    app.mainloop()
